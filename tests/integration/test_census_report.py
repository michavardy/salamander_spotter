"""Spec §9.4 — season census report (XLSX + PDF + CSV)."""

from __future__ import annotations

import pytest

from app.services import batches, census_report, ingest


@pytest.fixture
def published(db, settings, dataset_dir):
    ingest.transfer_dataset(db, settings, dataset_dir)
    return db


def test_xlsx_has_expected_sheets(published, settings):
    path = census_report.write_xlsx(published, settings.exports_dir)
    assert path.exists() and path.suffix == ".xlsx"
    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert set(wb.sheetnames) == {"Summary", "Roster", "Flagged"}
    assert wb["Roster"].max_row >= 2  # header + at least one published individual


def test_pdf_is_written(published, settings):
    path = census_report.write_pdf(published, settings.exports_dir)
    assert path.exists() and path.stat().st_size > 500
    assert path.read_bytes()[:4] == b"%PDF"


def test_roster_csv(published, settings):
    path = census_report.write_roster_csv(published, settings.exports_dir)
    text = path.read_text(encoding="utf-8")
    assert "display_id" in text.splitlines()[0]


def test_payload_reports_active_model(published):
    payload = census_report.report_payload(published)
    assert "stats" in payload and "roster" in payload
    assert payload["methods"]["model"] == "none"
